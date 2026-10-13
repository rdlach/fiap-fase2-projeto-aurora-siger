"""Gera relatorio_mgpeb.docx, a versao editavel em Word do relatorio.

Usa o mesmo texto (relatorio_mgpeb.md) e as mesmas medidas do PDF gerado por
gerar_pdf_relatorio.py: fontes DejaVu (embutidas no arquivo, para aparecerem
iguais em qualquer computador), tamanhos, entrelinhas, espacos, cores, larguras
das tabelas e paginacao automatica. Uso:

    python3 gerar_docx_relatorio.py

Requer python-docx, reportlab (so para medir as larguras) e fonttools.
"""

import io
import os
import re
import uuid
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION  # noqa: F401  (mantido para quem for editar)
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
import logging

from fontTools import subset
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

PASTA = Path(__file__).parent
ENTRADA_MD = PASTA / "relatorio_mgpeb.md"
SAIDA = PASTA / "relatorio_mgpeb.docx"
PASTA_FONTES = Path(os.environ.get("MGPEB_FONTES", "/usr/share/fonts/truetype/dejavu"))
if not (PASTA_FONTES / "DejaVuSans.ttf").exists():
    import matplotlib
    PASTA_FONTES = Path(matplotlib.get_data_path()) / "fonts" / "ttf"
FONTE, MONO = "DejaVu Sans", "DejaVu Sans Mono"
LARGURA_UTIL = Cm(17)
LARGURA_PT = 17 / 2.54 * 72
DESTAQUE = "D20A5C"
RODAPE = "Projeto Aurora Siger - Relatório Técnico do MGPEB - Fase 2"
# O PDF usa 7,8 pt nas tabelas e no codigo; o Word so aceita tamanhos em meio ponto.
TAM_TABELA = 8

# Quebras opcionais: (tipo do bloco, inicio do texto).
# "corte" divide um paragrafo antes do trecho indicado.
QUEBRAS = [("secao", "Referências")]  # Mantem o titulo junto da bibliografia.


# Fontes do PDF, registradas so para medir as larguras das colunas com a mesma regra.
pdfmetrics.registerFont(TTFont("Texto", str(PASTA_FONTES / "DejaVuSans.ttf")))
pdfmetrics.registerFont(TTFont("Texto-Negrito", str(PASTA_FONTES / "DejaVuSans-Bold.ttf")))


# ---------------------------------------------------------------- utilitarios
def cor(hexa):
    return RGBColor.from_string(hexa)


def fonte_run(run, nome=FONTE):
    run.font.name = nome
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), nome)


def formato(par, tamanho_linha, antes=0, depois=0, alinhamento=None, manter=False):
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(tamanho_linha)
    pf.space_before = Pt(antes)
    pf.space_after = Pt(depois)
    pf.widow_control = False
    if alinhamento is not None:
        par.alignment = alinhamento
    if manter:
        pf.keep_with_next = True


def espaco(doc, altura_pt, quebra=False):
    """Paragrafo vazio de altura exata (equivale ao Spacer do reportlab)."""
    p = doc.add_paragraph()
    formato(p, altura_pt)
    p.paragraph_format.page_break_before = quebra
    r = p.add_run()
    r.font.size = Pt(1)
    return p


def sombrear(celula, hexa):
    tcpr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexa)
    tcpr.append(shd)


def bordas_tabela(tabela, cor_hex, espessura_oitavos):
    tblpr = tabela._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(espessura_oitavos))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), cor_hex)
        b.append(e)
    tblpr.append(b)


def margens_celulas(tabela, esq, dir_, cima, baixo):
    tblpr = tabela._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for lado, v in (("top", cima), ("left", esq), ("bottom", baixo), ("right", dir_)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(int(v * 20)))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblpr.append(m)


def larguras_fixas(tabela, larguras_pt):
    tbl = tabela._tbl
    tblpr = tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "0")
    ind.set(qn("w:type"), "dxa")
    tblpr.append(ind)
    tabela.autofit = False
    tabela.alignment = WD_TABLE_ALIGNMENT.LEFT
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), larguras_pt):
        gc.set(qn("w:w"), str(int(round(w * 20))))
    for linha in tabela.rows:
        for celula, w in zip(linha.cells, larguras_pt):
            celula.width = Pt(w)


ORDEM_TBLPR = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
               "tblCaption", "tblDescription"]


def ordenar_tblpr(tabela):
    """Coloca as propriedades da tabela na ordem exigida pelo formato do Word."""
    tblpr = tabela._tbl.tblPr
    filhos = list(tblpr)
    for f in filhos:
        tblpr.remove(f)
    filhos.sort(key=lambda e: ORDEM_TBLPR.index(e.tag.split("}")[1]) if e.tag.split("}")[1] in ORDEM_TBLPR else 99)
    for f in filhos:
        tblpr.append(f)


def repetir_cabecalho(linha):
    trpr = linha._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    trpr.append(e)
    e2 = OxmlElement("w:cantSplit")
    trpr.append(e2)


# ---------------------------------------------------------------- texto de linha
PADRAO = re.compile(r"(`[^`]+`|\*\*.+?\*\*|(?<![\w*])\*(?!\s).+?(?<!\s)\*(?![\w*])|https?://\S+)")


def hyperlink(par, url, texto, tamanho, cor_hex):
    rid = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                             is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONTE)
    rpr.append(rf)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), cor_hex)
    rpr.append(c)
    for tag in ("w:sz", "w:szCs"):
        s = OxmlElement(tag)
        s.set(qn("w:val"), str(int(tamanho * 2)))
        rpr.append(s)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    par._p.append(h)


def escrever(par, texto, tamanho, cor_hex="000000", negrito=False, italico=False):
    pos = 0
    for m in PADRAO.finditer(texto):
        if m.start() > pos:
            _run(par, texto[pos:m.start()], tamanho, cor_hex, negrito, italico)
        t = m.group(0)
        if t.startswith("`"):
            _run(par, t[1:-1], tamanho, cor_hex, negrito, italico, MONO)
        elif t.startswith("**"):
            escrever(par, t[2:-2], tamanho, cor_hex, True, italico)
        elif t.startswith("*"):
            escrever(par, t[1:-1], tamanho, cor_hex, negrito, True)
        else:
            hyperlink(par, t, t, tamanho, DESTAQUE)
        pos = m.end()
    if pos < len(texto):
        _run(par, texto[pos:], tamanho, cor_hex, negrito, italico)


def _run(par, texto, tamanho, cor_hex, negrito, italico, nome=FONTE):
    r = par.add_run(texto)
    fonte_run(r, nome)
    r.font.size = Pt(tamanho)
    r.font.color.rgb = cor(cor_hex)
    r.bold = negrito
    r.italic = italico
    return r


# ---------------------------------------------------------------- blocos
def tabela(doc, linhas):
    celulas = [[c.strip() for c in l.strip().strip("|").split("|")] for l in linhas]
    celulas = [l for l in celulas if not all(re.fullmatch(r":?-{3,}:?", c) for c in l)]
    n = len(celulas[0])
    # mesma regra do PDF: minimo = maior palavra; o resto, proporcional ao texto
    minimos, pesos = [], []
    for col in range(n):
        textos = [l[col] for l in celulas]
        maior = 0
        for i, t in enumerate(textos):
            f = "Texto-Negrito" if i == 0 else "Texto"
            for p in re.sub(r"[`*]", "", t).split():
                maior = max(maior, pdfmetrics.stringWidth(p, f, 7.8))
        minimos.append(maior + 10)
        pesos.append(max(len(t) for t in textos))
    sobra = LARGURA_PT - sum(minimos)
    if sobra > 0:
        larguras = [m + sobra * p / sum(pesos) for m, p in zip(minimos, pesos)]
    else:
        larguras = [LARGURA_PT * m / sum(minimos) for m in minimos]

    t = doc.add_table(rows=len(celulas), cols=n)
    larguras_fixas(t, larguras)
    bordas_tabela(t, "999999", 3)
    margens_celulas(t, 4, 4, 3, 3)
    repetir_cabecalho(t.rows[0])
    ordenar_tblpr(t)
    for i, linha in enumerate(celulas):
        for j, texto in enumerate(linha):
            c = t.cell(i, j)
            p = c.paragraphs[0]
            formato(p, 9.8)
            if i == 0:
                sombrear(c, "232323")
                escrever(p, texto, TAM_TABELA, "FFFFFF", negrito=True)
            else:
                sombrear(c, "FFFFFF" if i % 2 == 1 else "F5F5F5")
                escrever(p, texto, TAM_TABELA)
    espaco(doc, 0.3 / 2.54 * 72)


def codigo(doc, linhas, quebra=False):
    espaco(doc, 0.1 / 2.54 * 72, quebra)
    t = doc.add_table(rows=1, cols=1)
    larguras_fixas(t, [LARGURA_PT])
    bordas_tabela(t, "CCCCCC", 4)
    margens_celulas(t, 8, 6, 6, 6)
    ordenar_tblpr(t)
    c = t.cell(0, 0)
    sombrear(c, "F2F2F2")
    for k, l in enumerate(linhas):
        p = c.paragraphs[0] if k == 0 else c.add_paragraph()
        formato(p, 10)
        _run(p, l if l else " ", TAM_TABELA, "000000", False, False, MONO)
    espaco(doc, 0.3 / 2.54 * 72)


def imagem(doc, legenda, caminho):
    arq = (PASTA / caminho).resolve()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(arq), width=LARGURA_UTIL)
    cap = doc.add_paragraph()
    formato(cap, 10.5, 3, 10, WD_ALIGN_PARAGRAPH.CENTER)
    escrever(cap, legenda, 8, "555555", italico=True)


# ---------------------------------------------------------------- documento
def preparar(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin = s.bottom_margin = Cm(1.8)
    s.header_distance = Cm(1)
    s.footer_distance = Cm(1.02)
    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(9)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONTE)
    normal.paragraph_format.space_after = Pt(0)
    # rodape: texto a esquerda e numero da pagina a direita
    f = s.footer.paragraphs[0]
    formato(f, 9)
    tabs = f.paragraph_format.tab_stops
    # o estilo de rodape do modelo traz tabulacoes proprias: limpa e usa so a da direita
    for pos in (4680, 9360):
        tabs.add_tab_stop(Pt(pos / 20), WD_TAB_ALIGNMENT.CLEAR)
    tabs.add_tab_stop(LARGURA_UTIL, WD_TAB_ALIGNMENT.RIGHT)
    _run(f, RODAPE + "\tPágina ", 7.5, "555555", False, False)
    for tipo in ("begin", "instr", "separate", "texto", "end"):
        r = _run(f, "", 7.5, "555555", False, False)
        if tipo == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = " PAGE "
            r._r.append(it)
        elif tipo == "texto":
            t = OxmlElement("w:t")
            t.text = "1"
            r._r.append(t)
        else:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tipo)
            r._r.append(e)


def quebra_para(tipo, texto):
    return any(t == tipo and texto.startswith(ini) for t, ini in QUEBRAS)


def converter(doc, texto):
    linhas = texto.splitlines()
    i = 0
    antes_da_primeira_secao = True
    while i < len(linhas):
        linha = linhas[i]
        if not linha.strip():
            i += 1
            continue
        if linha.startswith("```"):
            bloco = []
            i += 1
            while i < len(linhas) and not linhas[i].startswith("```"):
                bloco.append(linhas[i])
                i += 1
            i += 1
            codigo(doc, bloco, quebra_para("codigo", bloco[0].strip()))
            continue
        if linha.startswith("|"):
            bloco = []
            while i < len(linhas) and linhas[i].startswith("|"):
                bloco.append(linhas[i])
                i += 1
            tabela(doc, bloco)
            continue
        m = re.fullmatch(r"!\[(.*)\]\((.+)\)", linha.strip())
        if m:
            imagem(doc, m.group(1), m.group(2))
            i += 1
            continue
        if linha.startswith("# "):
            p = doc.add_paragraph(style="Title")
            limpar_estilo_titulo(p)
            formato(p, 23, 0, 8)
            escrever(p, linha[2:], 19, "202020", negrito=True)
            i += 1
            continue
        if linha.startswith("## "):
            antes_da_primeira_secao = False
            p = doc.add_paragraph(style="Heading 1")
            limpar_estilo_titulo(p)
            q = quebra_para("secao", linha[3:])
            formato(p, 16, 0 if q else 12, 6, manter=True)
            p.paragraph_format.page_break_before = q
            escrever(p, linha[3:], 13, DESTAQUE, negrito=True)
            i += 1
            continue
        if linha.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            limpar_estilo_titulo(p)
            q = quebra_para("subsecao", linha[4:])
            formato(p, 13, 0 if q else 6, 4, manter=True)
            p.paragraph_format.page_break_before = q
            escrever(p, linha[4:], 10.5, "333333", negrito=True)
            i += 1
            continue
        if linha.startswith("- "):
            itens = []
            while i < len(linhas) and linhas[i].startswith("- "):
                itens.append(linhas[i][2:])
                i += 1
            for k, item in enumerate(itens):
                p = doc.add_paragraph()
                ultimo = k == len(itens) - 1
                formato(p, 12.6, 0, 2 + (0.15 / 2.54 * 72 if ultimo else 0), WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.left_indent = Pt(12)
                p.paragraph_format.first_line_indent = Pt(-10)
                p.paragraph_format.tab_stops.add_tab_stop(Pt(12))
                _run(p, "•\t", 9, "000000", False, False)
                escrever(p, item, 9)
            continue
        par = []
        while i < len(linhas) and linhas[i].strip() and not re.match(r"(#|\||```|- |!\[)", linhas[i]):
            par.append(linhas[i].strip())
            i += 1
        texto_par = " ".join(par)
        if antes_da_primeira_secao:
            p = doc.add_paragraph()
            formato(p, 13, 0, 2, WD_ALIGN_PARAGRAPH.LEFT)
            escrever(p, texto_par, 9.5, "444444")
            continue
        partes = [texto_par]
        for t, trecho in QUEBRAS:
            if t == "corte" and trecho in texto_par:
                k = texto_par.index(trecho)
                partes = [texto_par[:k].rstrip(), texto_par[k:]]
        for k, parte in enumerate(partes):
            p = doc.add_paragraph()
            formato(p, 12.6, 0, 5 if k == len(partes) - 1 else 0, WD_ALIGN_PARAGRAPH.JUSTIFY)
            if k > 0 or quebra_para("paragrafo", parte):
                p.paragraph_format.page_break_before = True
            if len(partes) > 1 and k == 0:
                # a ultima linha antes do corte fica justificada, como no PDF
                justificar_ultima_linha(p)
            escrever(p, parte, 9)


def limpar_estilo_titulo(p):
    """Usa os estilos de titulo do Word (para o sumario e a navegacao), sem a aparencia padrao deles."""
    st = p.style
    st.font.name = FONTE
    st.font.color.rgb = None
    st.font.italic = False
    st.font.underline = False
    pPr = st.element.get_or_add_pPr()
    for e in pPr.findall(qn("w:pBdr")):
        pPr.remove(e)


def justificar_ultima_linha(p):
    jc = p._p.get_or_add_pPr().find(qn("w:jc"))
    jc.set(qn("w:val"), "distribute")


# ---------------------------------------------------------------- fontes embutidas
FAIXAS = ("U+0020-007E,U+00A0-017F,U+0391-03C9,U+2010-2027,U+2030-203A,U+2070-209F,"
          "U+2190-21FF,U+2200-22FF,U+25A0-25FF")

EMBUTIR = [
    (FONTE, {"Regular": "DejaVuSans.ttf", "Bold": "DejaVuSans-Bold.ttf",
             "Italic": "DejaVuSans-Oblique.ttf", "BoldItalic": "DejaVuSans-BoldOblique.ttf"}),
    (MONO, {"Regular": "DejaVuSansMono.ttf", "Bold": "DejaVuSansMono-Bold.ttf"}),
]


def subconjunto(arq):
    logging.getLogger("fontTools.subset").setLevel(logging.ERROR)
    opts = subset.Options()
    opts.name_IDs = ["*"]
    opts.notdef_outline = True
    opts.glyph_names = True
    opts.layout_features = ["*"]
    f = subset.load_font(str(arq), opts)
    s = subset.Subsetter(opts)
    s.populate(unicodes=subset.parse_unicodes(FAIXAS))
    s.subset(f)
    buf = io.BytesIO()
    f.save(buf)
    return buf.getvalue()


def ofuscar(dados, chave):
    # ECMA-376: XOR dos 32 primeiros bytes com a chave GUID em ordem inversa
    hexa = chave.strip("{}").replace("-", "")
    k = bytes(int(hexa[i:i + 2], 16) for i in range(30, -1, -2))
    b = bytearray(dados)
    for i in range(32):
        b[i] ^= k[i % 16]
    return bytes(b)


def embutir_fontes(caminho):
    with zipfile.ZipFile(caminho) as z:
        arquivos = {n: z.read(n) for n in z.namelist()}
    fontes_xml, rels, n = [], [], 0
    for nome, estilos in EMBUTIR:
        refs = ""
        for estilo, arq in estilos.items():
            n += 1
            chave = "{" + str(uuid.uuid4()).upper() + "}"
            arquivos[f"word/fonts/font{n}.odttf"] = ofuscar(subconjunto(PASTA_FONTES / arq), chave)
            rels.append(f'<Relationship Id="rIdFonte{n}" Type="http://schemas.openxmlformats.org/officeDocument/'
                        f'2006/relationships/font" Target="fonts/font{n}.odttf"/>')
            refs += f'<w:embed{estilo} r:id="rIdFonte{n}" w:fontKey="{chave}"/>'
        fontes_xml.append(f'<w:font w:name="{nome}"><w:charset w:val="00"/><w:family w:val="swiss"/>'
                          f'<w:pitch w:val="{"fixed" if "Mono" in nome else "variable"}"/>{refs}</w:font>')
    ft = arquivos["word/fontTable.xml"].decode("utf-8")
    for nome, _ in EMBUTIR:
        ft = re.sub(rf'<w:font w:name="{re.escape(nome)}">.*?</w:font>', "", ft, flags=re.S)
    if 'xmlns:r=' not in ft.split(">", 2)[1]:
        ft = re.sub(r"<w:fonts ", '<w:fonts xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
                                  'relationships" ', ft, count=1)
    ft = ft.replace("</w:fonts>", "".join(fontes_xml) + "</w:fonts>")
    arquivos["word/fontTable.xml"] = ft.encode("utf-8")
    arquivos["word/_rels/fontTable.xml.rels"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(rels) + "</Relationships>").encode("utf-8")
    ct = arquivos["[Content_Types].xml"].decode("utf-8")
    if 'Extension="odttf"' not in ct:
        ct = ct.replace("<Default ", '<Default Extension="odttf" ContentType="application/vnd.openxmlformats-'
                                     'officedocument.obfuscatedFont"/><Default ', 1)
    arquivos["[Content_Types].xml"] = ct.encode("utf-8")
    st = arquivos["word/settings.xml"].decode("utf-8")
    st = re.sub(r"<w:embedTrueTypeFonts/>|<w:saveSubsetFonts/>", "", st)
    # o modelo do python-docx traz o zoom sem o atributo obrigatorio
    st = re.sub(r"<w:zoom (?![^>]*w:percent)", '<w:zoom w:percent="100" ', st)
    if "<w:zoom" in st:
        st = re.sub(r"(<w:zoom[^>]*/>)", r"\1<w:embedTrueTypeFonts/><w:saveSubsetFonts/>", st, count=1)
    else:
        st = re.sub(r"(<w:settings[^>]*>)", r"\1<w:embedTrueTypeFonts/><w:saveSubsetFonts/>", st, count=1)
    # modo de compatibilidade do Word 2013+ (tabelas alinhadas a margem)
    st = re.sub(r'<w:compatSetting w:name="compatibilityMode"[^>]*/>', "", st)
    modo = ('<w:compatSetting w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" '
            'w:val="15"/>')
    st = st.replace("</w:compat>", modo + "</w:compat>") if "</w:compat>" in st else \
        st.replace("</w:settings>", f"<w:compat>{modo}</w:compat></w:settings>")
    arquivos["word/settings.xml"] = st.encode("utf-8")
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as z:
        for nome, dados in arquivos.items():
            z.writestr(nome, dados)


def gerar():
    doc = Document()
    preparar(doc)
    doc.core_properties.title = "Relatório Técnico - MGPEB - Aurora Siger"
    doc.core_properties.author = "Grupo 42"
    converter(doc, ENTRADA_MD.read_text(encoding="utf-8"))
    # o modelo do python-docx comeca com um paragrafo vazio: remove
    corpo = doc.element.body
    primeiro = corpo[0]
    if primeiro.tag == qn("w:p") and not "".join(primeiro.itertext()).strip():
        corpo.remove(primeiro)
    doc.save(SAIDA)
    embutir_fontes(SAIDA)
    print(f"DOCX gerado em {SAIDA}")


if __name__ == "__main__":
    gerar()
